"""Exercise five real snapshot report shapes through the API and worker (development evidence only)."""

import json
import os
import time
import uuid
from pathlib import Path

import httpx
from docx import Document


def main():
    output = Path("artifacts/reports")
    output.mkdir(parents=True, exist_ok=True)
    scenarios = {
        "summary": "2026 年 9 月销售额是多少？",
        "empty": "2020 年 1 月销售额是多少？",
        "share": "2026 年 9 月各类别销售额占比，饼图",
        "trend": "2026 年各商品类别的月度销售额趋势",
        "detail1000": "列出全部历史订单明细",
    }
    evidence = []
    with httpx.Client(base_url="http://127.0.0.1:8000", timeout=20) as client:
        login = client.post(
            "/api/auth/login",
            json={"username": "demo", "password": os.environ.get("ERP_EVAL_PASSWORD", "erp-demo-2026")},
        )
        login.raise_for_status()
        client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
        for name, question in scenarios.items():
            conv = client.post("/api/conversations", json={"title": "报告验证 " + name}).json()["id"]
            response = client.post(
                f"/api/conversations/{conv}/queries",
                json={"question": question, "idempotency_key": str(uuid.uuid4())},
            )
            response.raise_for_status()
            result = response.json()
            assert result["status"] == "success", result
            all_rows = []
            page = 1
            while len(all_rows) < result["total_rows"]:
                all_rows.extend(client.get(f"/api/results/{result['result_id']}?page={page}").json()["rows"])
                page += 1
            expected = [
                ["—" if row.get(c["key"]) is None else str(row[c["key"]]) for c in result["columns"]]
                for row in all_rows
            ]
            for format in ("docx", "pdf"):
                started = time.perf_counter()
                response = client.post(
                    f"/api/results/{result['result_id']}/exports",
                    json={"format": format, "include_sql": True},
                )
                response.raise_for_status()
                job = response.json()["id"]
                deadline = time.monotonic() + 180
                while time.monotonic() < deadline:
                    state = client.get(f"/api/exports/{job}").json()
                    assert state["status"] not in {"failed", "expired"}, state
                    if state["status"] == "succeeded":
                        break
                    time.sleep(0.25)
                else:
                    raise RuntimeError("export deadline exceeded")
                download = client.get(f"/api/exports/{job}/download")
                download.raise_for_status()
                path = output / f"{name}.{format}"
                path.write_bytes(download.content)
                if format == "docx":
                    doc = Document(path)
                    actual = [[c.text for c in row.cells] for row in doc.tables[0].rows[1:]]
                    assert actual == expected, name
                    assert result["answer"] in [p.text for p in doc.paragraphs]
                    assert any("SQL 附录" in p.text for p in doc.paragraphs)
                else:
                    assert download.content.startswith(b"%PDF-")
                evidence.append(
                    {
                        "scenario": name,
                        "format": format,
                        "saved_rows": len(all_rows),
                        "truncated": result["truncated"],
                        "elapsed_seconds": round(time.perf_counter() - started, 3),
                        "checksum": state["checksum"],
                    }
                )
    (output / "verification.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("10 reports generated; all 5 Word tables match their entire saved snapshot exactly.")


if __name__ == "__main__":
    main()
