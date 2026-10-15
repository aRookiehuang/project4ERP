import base64
import hashlib
import html
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from playwright.sync_api import sync_playwright

from backend.browser import browser_launch_options
from .config import Settings


def display(value):
    return "—" if value is None else str(value)


def report_html(payload, include_sql, font_css=""):
    esc = lambda x: html.escape(display(x))
    columns = payload["columns"]
    headers = "".join(
        f"<th>{esc(c['label'])}{'（' + esc(c['unit']) + '）' if c['unit'] else ''}</th>" for c in columns
    )
    rows = "".join(
        "<tr>" + "".join(f"<td>{esc(row.get(c['key']))}</td>" for c in columns) + "</tr>"
        for row in payload["rows"]
    )
    definitions = "；".join(
        f"{v['label']}：{v['definition']}" for v in payload["metric_definitions"].values()
    )
    return f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><style>
    {font_css}
    body{{font-family:'Noto Sans CJK SC',sans-serif;color:#172b38;font-size:11px;line-height:1.65}}
    h1{{font-size:25px;color:#087f72}} .meta{{color:#526775;overflow-wrap:anywhere}}
    table{{border-collapse:collapse;width:100%;table-layout:fixed}}th,td{{border:1px solid #d9e2e7;padding:6px;overflow-wrap:anywhere}}
    th{{background:#eef5f4}}thead{{display:table-header-group}}tr{{break-inside:avoid}}pre{{white-space:pre-wrap;overflow-wrap:anywhere}}
    #chart{{width:900px;height:380px;max-width:100%;break-inside:avoid}}h2{{font-size:16px}}
    </style><body><h1>ERP 销售分析报告</h1><h2>{esc(payload["question"])}</h2>
    <p>{esc(payload["answer"])}</p><p class="meta">时间：{esc(payload["time_range"]["start"])} ≤ 时间 &lt; {esc(payload["time_range"]["end"])}（Asia/Shanghai）<br>
    {esc(" ".join(payload["assumptions"]))}<br>{esc(definitions)}<br>
    数据覆盖：{esc(" 至 ".join(payload["data_coverage"]))}；数据版本：{esc(payload["versions"]["data"])}<br>
    快照时间：{esc(payload["snapshot_at"])}；追溯 ID：{esc(payload["query_id"])}</p>
    {"<p>结果已截断，仅保留前 1,000 行。</p>" if payload["truncated"] else ""}
    <div id="chart"></div><table><thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table>
    {"<h2>SQL 附录（参数隐藏）</h2><pre>" + esc(payload["sql_display"]) + "</pre>" if include_sql else ""}</body></html>"""


CHART_RENDER = """({spec, columns}) => {
 const node = document.querySelector('#chart');
 if (!['bar','line','pie'].includes(spec.type)) { node.remove(); return; }
 const labels = Object.fromEntries(columns.map(c => [c.key, c.label + (c.unit ? '（'+c.unit+'）' : '')]));
 const units = [...new Set(spec.series.map(m => columns.find(c => c.key === m)?.unit || ''))];
 const key = (row, dimension) => row[dimension+'_id'] == null ? String(row[dimension]) : row[dimension]+' #'+row[dimension+'_id'];
 const x = [...new Set(spec.data.map(r => key(r, spec.x)))];
 const chart = echarts.init(node, null, {renderer:'svg'});
 let series=[];
 for(const metric of spec.series){
  const groups = spec.group ? [...new Set(spec.data.map(r => key(r, spec.group)))] : [''];
  for(const group of groups){
   const data = x.map(label => {
    const row=spec.data.find(r=>key(r,spec.x)===label && (!spec.group || key(r,spec.group)===group));
    return row && row[metric] != null ? Number(row[metric]) : null;
   });
   series.push({name:(group ? group+' · ' : '')+labels[metric],type:spec.type,data,connectNulls:false,
     yAxisIndex:units.indexOf(columns.find(c=>c.key===metric)?.unit || '')});
  }
 }
 const option={animation:false,textStyle:{fontFamily:'Noto Sans CJK SC'},color:['#0d9488','#5b8def','#eba94e','#a78bdb','#e58193'],
  legend:{type:'scroll',bottom:0},grid:{left:65,right:units.length>1?75:25,top:35,bottom:80,containLabel:true},
  xAxis:{type:'category',data:x,axisLabel:{rotate:x.length>8?30:0}},
  yAxis:units.map((unit,index)=>({type:'value',name:unit,position:index?'right':'left',offset:index>1?50:0})),series};
 if(spec.type==='pie'){delete option.xAxis;delete option.yAxis;option.series=[{type:'pie',radius:['35%','65%'],data:spec.data.map(r=>({name:key(r,spec.x),value:Number(r[spec.series[0]])}))}];}
 chart.setOption(option);
} """


def render_report(payload: dict, format: str, include_sql: bool, output: Path, config: Settings):
    font_css = ""
    if config.font_path.is_file():
        encoded = base64.b64encode(config.font_path.read_bytes()).decode()
        font_css = (
            "@font-face{font-family:'Noto Sans CJK SC';src:url(data:font/otf;base64,"
            + encoded
            + ") format('opentype')}"
        )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            **browser_launch_options(
                channel=config.browser_channel, executable_path=config.browser_executable_path
            )
        )
        printable_width = round(((297 if len(payload["columns"]) > 7 else 210) - 24) * 96 / 25.4)
        context = browser.new_context(viewport={"width": printable_width, "height": 1000})
        context.route("**/*", lambda route: route.abort())
        page = context.new_page()
        page.set_default_timeout(15_000)
        page.set_content(report_html(payload, include_sql, font_css), wait_until="load")
        page.evaluate("document.fonts.ready")
        if payload["chart_spec"]["type"] in {"bar", "line", "pie"}:
            if not config.chart_js_path.is_file():
                raise RuntimeError("ECharts local asset missing")
            page.add_script_tag(path=str(config.chart_js_path))
        page.evaluate(CHART_RENDER, {"spec": payload["chart_spec"], "columns": payload["columns"]})
        if format == "pdf":
            page.pdf(
                path=str(output),
                format="A4",
                landscape=len(payload["columns"]) > 7,
                print_background=True,
                margin={"top": "16mm", "bottom": "18mm", "left": "12mm", "right": "12mm"},
                display_header_footer=True,
                header_template="<span></span>",
                footer_template='<div style="font-size:8px;width:100%;text-align:center">ERP · <span class="pageNumber"></span> / <span class="totalPages"></span></div>',
            )
        else:
            chart_png = page.locator("#chart").screenshot() if page.locator("#chart").count() else None
            document = Document()
            style = document.styles["Normal"]
            style.font.name = "Noto Sans CJK SC"
            style.font.size = Pt(10)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
            if len(payload["columns"]) > 7:
                section = document.sections[0]
                section.page_width, section.page_height = section.page_height, section.page_width
            document.add_heading("ERP 销售分析报告", 0)
            document.add_heading(payload["question"], 1)
            document.add_paragraph(payload["answer"])
            document.add_paragraph(
                f"时间：{payload['time_range']['start']} ≤ 时间 < {payload['time_range']['end']}（Asia/Shanghai）"
            )
            for assumption in payload["assumptions"]:
                document.add_paragraph(assumption)
            for metric in payload["metric_definitions"].values():
                document.add_paragraph(metric["label"] + "：" + metric["definition"])
            document.add_paragraph(
                "数据覆盖："
                + " 至 ".join(payload["data_coverage"])
                + "；版本："
                + payload["versions"]["data"]
            )
            document.add_paragraph(f"快照时间：{payload['snapshot_at']}；追溯 ID：{payload['query_id']}")
            if payload["truncated"]:
                document.add_paragraph("结果已截断，仅保留前 1,000 行。")
            if chart_png:
                document.add_picture(BytesIO(chart_png), width=Inches(6))
            table = document.add_table(rows=1, cols=len(payload["columns"]))
            table.style = "Light Shading Accent 1"
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for cell, column in zip(table.rows[0].cells, payload["columns"]):
                cell.text = column["label"] + (f"（{column['unit']}）" if column["unit"] else "")
            for row in payload["rows"]:
                for cell, column in zip(table.add_row().cells, payload["columns"]):
                    cell.text = display(row.get(column["key"]))
            if include_sql:
                document.add_heading("SQL 附录（参数隐藏）", 1)
                document.add_paragraph(payload["sql_display"])
            document.save(str(output))
        context.close()
        browser.close()
    return hashlib.sha256(output.read_bytes()).hexdigest()
